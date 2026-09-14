"""解析器与上传校验测试：扩展名白名单、Content-Type、文件头、ZIP bomb、
Text/Pdf/Docx 提取。PDF 与 DOCX fixture 全部动态生成（零网络、无二进制依赖）。"""

import io
import zipfile

import pytest

from app.services.document_errors import (
    DocxZipBombError,
    InvalidFileContentError,
    UnsupportedFileTypeError,
)
from app.services.parsers import (
    DocxParser,
    ParsingError,
    PdfParser,
    TextParser,
    check_content_type,
    check_docx_zip_bomb,
    check_extension,
    check_file_header,
    get_parser,
)


# ---------- fixture 生成器 ----------


def _make_minimal_pdf(text: str = "Hello PDF") -> bytes:
    """手写最小单页 PDF：构造对象流、计算字节偏移、生成 xref 表。

    必须带 /Font 资源（pypdf 6.x 对无字体映射的 content stream 提取为空）。
    """
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >> endobj",
        b"4 0 obj << /Length " + str(len(stream)).encode() + b" >> stream\n"
        + stream + b"\nendstream endobj",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(out))
        out += obj + b"\n"
    xref_pos = len(out)
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
    out += str(xref_pos).encode() + b"\n%%EOF\n"
    return bytes(out)


def _make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                table.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_encrypted_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_zip_bomb_ratio() -> bytes:
    """重复字节 deflate 后压缩比远超 200:1。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.txt", b"A" * (5 * 1024 * 1024))
    return buf.getvalue()


def _make_zip_bomb_entries(count: int = 1001) -> bytes:
    """条目数超过上限（1001 > 1000）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(count):
            zf.writestr(f"f{i}.txt", b"x")
    return buf.getvalue()


# ---------- 扩展名白名单 ----------


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("report.txt", ".txt"),
        ("NOTES.md", ".md"),
        ("guide.pdf", ".pdf"),
        ("proposal.docx", ".docx"),
        ("archive.tar.gz", None),  # 扩展名不在白名单
        ("no_extension", None),
        ("evil.txt.exe", None),  # 只看最后扩展名
    ],
)
def test_check_extension(filename, expected):
    if expected is None:
        with pytest.raises(UnsupportedFileTypeError):
            check_extension(filename)
    else:
        assert check_extension(filename) == expected


def test_check_extension_rejects_docm_explicitly():
    with pytest.raises(UnsupportedFileTypeError) as excinfo:
        check_extension("macro.docm")
    assert "docm" in str(excinfo.value)


# ---------- Content-Type ----------


def test_check_content_type_matching():
    check_content_type("text/plain; charset=utf-8", ".txt")
    check_content_type("text/markdown", ".md")
    check_content_type("application/pdf", ".pdf")
    check_content_type(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"
    )


def test_check_content_type_neutral_types_pass():
    for neutral in ("", "application/octet-stream", "application/x-www-form-urlencoded"):
        check_content_type(neutral, ".pdf")  # 中性类型交给文件头兜底


def test_check_content_type_conflict_rejected():
    with pytest.raises(InvalidFileContentError):
        check_content_type("image/png", ".pdf")
    with pytest.raises(InvalidFileContentError):
        check_content_type("text/html", ".docx")


# ---------- 文件头（伪装检测） ----------


def test_check_file_header_matches():
    check_file_header(b"%PDF-1.4 ...", ".pdf")
    check_file_header(b"PK\x03\x04 rest", ".docx")
    check_file_header("任意文本".encode("utf-8"), ".txt")  # 无强制文件头


@pytest.mark.parametrize(
    "content, ext",
    [
        (b"plain text not pdf", ".pdf"),
        (b"%PDF-1.4 but named docx", ".docx"),  # 内容为 PDF，伪装成 docx
        (b"", ".pdf"),  # 空文件不能是 PDF
    ],
)
def test_check_file_header_mismatch_rejected(content, ext):
    with pytest.raises(InvalidFileContentError):
        check_file_header(content, ext)


# ---------- ZIP bomb 防御 ----------


def test_docx_zip_bomb_ratio_detected():
    with pytest.raises(DocxZipBombError):
        check_docx_zip_bomb(_make_zip_bomb_ratio())


def test_docx_zip_bomb_entries_detected():
    with pytest.raises(DocxZipBombError):
        check_docx_zip_bomb(_make_zip_bomb_entries())


def test_docx_zip_bomb_valid_docx_passes():
    check_docx_zip_bomb(_make_docx(["段落一", "段落二"]))


def test_docx_zip_bomb_corrupt_zip_rejected():
    with pytest.raises(InvalidFileContentError):
        check_docx_zip_bomb(b"PK\x03\x04 not a real zip archive at all")


# ---------- TextParser ----------


def test_text_parser_utf8():
    assert TextParser(".txt").extract_text("你好，世界".encode("utf-8")) == "你好，世界"


def test_text_parser_strips_bom():
    text = "with bom"
    assert TextParser(".md").extract_text(b"\xef\xbb\xbf" + text.encode("utf-8")) == text


def test_text_parser_replaces_invalid_bytes_deterministically():
    result = TextParser(".txt").extract_text(b"ab\xffcd")
    assert result == "ab�cd"


# ---------- PdfParser ----------


def test_pdf_parser_extracts_text():
    text = PdfParser().extract_text(_make_minimal_pdf("Hello PDF"))
    assert "Hello PDF" in text


def test_pdf_parser_rejects_encrypted():
    with pytest.raises(ParsingError):
        PdfParser().extract_text(_make_encrypted_pdf())


def test_pdf_parser_corrupt_pdf_raises():
    with pytest.raises(ParsingError):
        PdfParser().extract_text(b"%PDF-1.4 garbage not a real pdf")


# ---------- DocxParser ----------


def test_docx_parser_extracts_paragraphs_in_order():
    docx = _make_docx(["第一段", "第二段", "第三段"])
    text = DocxParser().extract_text(docx)
    lines = text.split("\n")
    assert lines == ["第一段", "第二段", "第三段"]


def test_docx_parser_extracts_tables_after_paragraphs():
    docx = _make_docx(["开头段"], table_rows=[["姓名", "年龄"], ["张三", "30"]])
    text = DocxParser().extract_text(docx)
    assert text.split("\n") == ["开头段", "姓名 | 年龄", "张三 | 30"]


def test_docx_parser_corrupt_docx_raises():
    with pytest.raises(ParsingError):
        DocxParser().extract_text(b"PK\x03\x04 truncated garbage")


# ---------- 解析器注册表 ----------


def test_get_parser_by_extension():
    assert isinstance(get_parser(".txt"), TextParser)
    assert isinstance(get_parser(".md"), TextParser)
    assert isinstance(get_parser(".pdf"), PdfParser)
    assert isinstance(get_parser(".docx"), DocxParser)
