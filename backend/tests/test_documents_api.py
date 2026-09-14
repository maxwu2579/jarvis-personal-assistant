"""文档 API 集成测试：上传（txt/md/pdf/docx）、校验错误码、列表/详情/
块预览/删除、状态机 FAILED、故障注入、磁盘一致性。"""

import io
import zipfile

import pytest

API = "/api/documents"


# ---------- 二进制 fixture ----------


def _make_minimal_pdf(text: str = "Hello PDF") -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >> endobj",
        b"4 0 obj << /Length " + str(len(stream)).encode() + b" >> stream\n"
        + stream + b"\nendstream endobj",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(out))
        out += obj + b"\n"
    xref_pos = len(out)
    out += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
    out += str(xref_pos).encode() + b"\n%%EOF\n"
    return bytes(out)


def _make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, cell in enumerate(row):
                table.rows[i].cells[j].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_encrypted_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_zip_bomb_ratio() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.txt", b"A" * (5 * 1024 * 1024))
    return buf.getvalue()


def _make_zip_bomb_entries(count: int = 1001) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(count):
            zf.writestr(f"f{i}.txt", b"x")
    return buf.getvalue()


def _upload(client, filename: str, content: bytes, content_type: str):
    return client.post(
        API, files={"file": (filename, content, content_type)}
    )


def _upload_txt(client, text: str, filename: str = "note.txt"):
    return _upload(client, filename, text.encode("utf-8"), "text/plain")


# ---------- 上传成功路径 ----------


def test_upload_txt_ready_with_chunks(client):
    text = "JARVIS 文档摄取测试。" * 60
    resp = _upload_txt(client, text)
    assert resp.status_code == 201
    body = resp.json()
    assert body["status"] == "READY"
    assert body["original_filename"] == "note.txt"
    assert body["size_bytes"] == len(text.encode("utf-8"))
    assert len(body["sha256"]) == 64
    assert body["error_code"] is None
    assert body["content_type"] == "text/plain"
    # 块可查且有序
    chunks = client.get(f"{API}/{body['id']}/chunks")
    assert chunks.status_code == 200
    chunk_list = chunks.json()
    assert len(chunk_list) >= 2
    assert [c["chunk_index"] for c in chunk_list] == list(range(len(chunk_list)))
    for chunk in chunk_list:
        assert chunk["document_id"] == body["id"]
        assert chunk["char_start"] < chunk["char_end"]
        assert chunk["token_estimate"] > 0


def test_upload_md_document(client):
    resp = _upload(client, "notes.md", b"# Title\n\nbody", "text/markdown")
    assert resp.status_code == 201
    assert resp.json()["status"] == "READY"


def test_upload_pdf_document(client):
    resp = _upload(client, "paper.pdf", _make_minimal_pdf("Hello PDF"), "application/pdf")
    assert resp.status_code == 201
    assert resp.json()["status"] == "READY"
    chunks = client.get(f"{API}/{resp.json()['id']}/chunks").json()
    assert "Hello PDF" in chunks[0]["content"]


def test_upload_docx_document_extracts_table(client):
    docx = _make_docx(["开头段"], table_rows=[["姓名", "年龄"], ["张三", "30"]])
    resp = _upload(
        client,
        "report.docx",
        docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert resp.status_code == 201
    assert resp.json()["status"] == "READY"
    chunks = client.get(f"{API}/{resp.json()['id']}/chunks").json()
    joined = "".join(c["content"] for c in chunks)
    assert "开头段" in joined
    assert "张三" in joined and "30" in joined


def test_upload_empty_txt_ready_with_no_chunks(client):
    resp = _upload_txt(client, "")
    assert resp.status_code == 201
    assert resp.json()["status"] == "READY"
    assert resp.json()["size_bytes"] == 0
    assert client.get(f"{API}/{resp.json()['id']}/chunks").json() == []


def test_same_content_uploads_create_two_documents(client):
    a = _upload_txt(client, "duplicate content")
    b = _upload_txt(client, "duplicate content")
    assert a.status_code == b.status_code == 201
    assert a.json()["sha256"] == b.json()["sha256"]
    assert a.json()["id"] != b.json()["id"]


# ---------- 上传校验错误码（稳定） ----------


def test_upload_docm_rejected_415(client):
    resp = _upload(client, "macro.docm", b"PK\x03\x04garbage", "application/octet-stream")
    assert resp.status_code == 415
    assert resp.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


def test_upload_unknown_extension_rejected_415(client):
    resp = _upload(client, "script.exe", b"MZ...", "application/octet-stream")
    assert resp.status_code == 415
    assert resp.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"
    assert "txt" in resp.json()["error"]["message"]  # 白名单可见


def test_upload_extension_masquerade_rejected_422(client):
    """扩展名说 pdf，内容是文本 → 422 INVALID_FILE_CONTENT（伪装检测）。"""
    resp = _upload(client, "fake.pdf", b"this is not a pdf", "application/pdf")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "INVALID_FILE_CONTENT"


def test_upload_docx_masquerade_rejected_422(client):
    resp = _upload(client, "fake.docx", b"also not a zip", "application/octet-stream")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "INVALID_FILE_CONTENT"


def test_upload_content_type_conflict_rejected_422(client):
    resp = _upload(client, "image.pdf", _make_minimal_pdf(), "image/png")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "INVALID_FILE_CONTENT"


def test_upload_corrupt_docx_zip_rejected_422(client):
    resp = _upload(client, "broken.docx", b"PK\x03\x04not a real zip", "application/octet-stream")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "INVALID_FILE_CONTENT"


def test_upload_zip_bomb_ratio_rejected_422(client):
    resp = _upload(client, "bomb.docx", _make_zip_bomb_ratio(), "application/octet-stream")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "DOCX_ZIP_BOMB_DETECTED"


def test_upload_zip_bomb_entries_rejected_422(client):
    resp = _upload(client, "many.docx", _make_zip_bomb_entries(), "application/octet-stream")
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "DOCX_ZIP_BOMB_DETECTED"


def test_upload_too_large_rejected_413(client, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "document_max_file_size_mb", 1)
    resp = _upload_txt(client, "x" * (1024 * 1024 + 1))
    assert resp.status_code == 413
    assert resp.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_upload_missing_filename_rejected_422(client):
    resp = client.post(API, files={"file": ("", b"content", "text/plain")})
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "VALIDATION_ERROR"


# ---------- 处理失败（FAILED 状态，201 返回） ----------


def test_encrypted_pdf_upload_returns_failed_document(client):
    resp = _upload(client, "secret.pdf", _make_encrypted_pdf(), "application/pdf")
    assert resp.status_code == 201
    body = resp.json()
    assert body["status"] == "FAILED"
    assert body["error_code"] == "DOCUMENT_PROCESSING_FAILED"
    assert body["error_message"]
    # 非 READY 文档取 chunks → 409
    chunks = client.get(f"{API}/{body['id']}/chunks")
    assert chunks.status_code == 409
    assert chunks.json()["error"]["code"] == "INVALID_DOCUMENT_STATE"


# ---------- 列表 / 详情 / 删除 ----------


def test_list_documents_newest_first(client):
    first = _upload_txt(client, "first").json()
    second = _upload_txt(client, "second").json()
    resp = client.get(API)
    assert resp.status_code == 200
    ids = [d["id"] for d in resp.json()]
    assert ids[:2] == [second["id"], first["id"]]


def test_get_document_detail(client):
    created = _upload_txt(client, "detail me").json()
    resp = client.get(f"{API}/{created['id']}")
    assert resp.status_code == 200
    assert resp.json()["original_filename"] == "note.txt"
    assert resp.json()["status"] == "READY"


def test_get_document_not_found_404(client):
    resp = client.get(f"{API}/9999")
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "DOCUMENT_NOT_FOUND"


def test_chunks_preview_truncated(client):
    """长块内容只返回 300 字符预览 + 完整长度标注。"""
    resp = _upload_txt(client, "A" * 5000, filename="long.txt")
    doc_id = resp.json()["id"]
    chunks = client.get(f"{API}/{doc_id}/chunks").json()
    assert len(chunks) >= 1
    first = chunks[0]
    assert first["content_length"] > 300
    assert len(first["content"]) == 301  # 300 + 省略号
    assert first["content"].endswith("…")
    # 短块不截断
    short = _upload_txt(client, "短文本").json()
    short_chunk = client.get(f"{API}/{short['id']}/chunks").json()[0]
    assert short_chunk["content"] == "短文本"
    assert short_chunk["content_length"] == 3


def test_delete_document_removes_row_file_and_chunks(client, _isolated_document_storage):
    doc = _upload_txt(client, ("句子。" * 200).encode("utf-8").decode("utf-8")).json()
    # storage_key 不通过 API 暴露：从磁盘目录定位存储文件（唯一文件）
    stored_files = list(_isolated_document_storage.iterdir())
    assert len(stored_files) == 1
    storage_file = stored_files[0]
    assert storage_file.name.isalnum() and storage_file.name.islower()  # UUID 形态
    resp = client.delete(f"{API}/{doc['id']}")
    assert resp.status_code == 200
    assert resp.json() == {"deleted": True}
    assert not storage_file.exists()
    assert client.get(f"{API}/{doc['id']}").status_code == 404
    # 磁盘目录无残留
    assert list(_isolated_document_storage.iterdir()) == []


def test_delete_document_not_found_404(client):
    resp = client.delete(f"{API}/9999")
    assert resp.status_code == 404
    assert resp.json()["error"]["code"] == "DOCUMENT_NOT_FOUND"


def test_error_responses_do_not_leak_internals(client):
    """所有 4xx 错误响应不包含堆栈、路径、文件内容。"""
    resp = _upload(client, "fake.pdf", b"not a pdf", "application/pdf")
    body = resp.text
    assert "Traceback" not in body
    assert "\\\\" not in body
    assert "C:" not in body
    assert "not a pdf" not in body
