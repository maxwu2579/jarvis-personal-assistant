"""结构检查：重新打开生成的 DOCX，验证排版规则是否满足。

检查项：
1. 中英段落配对（每个英文段落在其前面存在中文内容段落）；
2. 真实 Heading 层级数量；
3. 表格数量与单元格非空；
4. 代码块存在且带浅色底纹；
5. 全文无 Unicode 拼图字符（box-drawing）；
6. 无 U+FFFD 替换字符（乱码标志）。

用法：python verify_docs.py [docs/learning/*.docx ...]（默认检查全部）
"""

import glob as _glob
import sys
from pathlib import Path

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_GLOB = PROJECT_ROOT / "docs" / "learning" / "*.docx"

BOX_DRAWING = set("─│┌┐└┘├┤┬┴┼╔╗╚╝╠╣╦╩╬═║")

CN_STYLES = {"CnPara", "Normal", "CellCn", "FieldLabel"}
EN_STYLE = "EnPara"
CODE_STYLE = "CodeBlock"


def check_document(path: Path) -> list[str]:
    problems: list[str] = []
    doc = Document(str(path))
    paragraphs = doc.paragraphs
    text_all = "\n".join(p.text for p in paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_all += "\n" + cell.text

    # 1. 中英配对：EnPara 前 2 段内应有中文内容段落
    en_count = 0
    unpaired = 0
    for i, p in enumerate(paragraphs):
        if p.style.name == EN_STYLE:
            en_count += 1
            window = paragraphs[max(0, i - 2):i]
            if not any(q.style.name in CN_STYLES and q.text.strip() for q in window):
                unpaired += 1
    if en_count == 0:
        problems.append("没有发现英文段落（EnPara 样式）")
    if unpaired:
        problems.append(f"{unpaired} 个英文段落没有找到前邻中文段落")

    # 2. Heading 层级
    headings = {"Heading 1": 0, "Heading 2": 0, "Heading 3": 0}
    for p in paragraphs:
        if p.style.name in headings:
            headings[p.style.name] += 1
    for level, count in headings.items():
        if count == 0:
            problems.append(f"缺少 {level} 标题")

    # 3. 表格
    if not doc.tables:
        problems.append("没有表格")
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                if not cell.text.strip():
                    problems.append(f"发现空单元格（表 {doc.tables.index(t) + 1}）")
                    break
            else:
                continue
            break

    # 4. 代码块与底纹
    code_paras = [p for p in paragraphs if p.style.name == CODE_STYLE]
    if not code_paras:
        problems.append("没有代码块")
    else:
        shaded = sum(1 for p in code_paras if p._p.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd") is not None)
        if shaded < len(code_paras):
            problems.append(f"代码块底纹缺失：{len(code_paras) - shaded} 段未加底纹")

    # 5. Unicode 拼图字符
    found_box = sorted({ch for ch in text_all if ch in BOX_DRAWING})
    if found_box:
        problems.append(f"发现 box-drawing 字符（疑似 Unicode 拼图）：{''.join(found_box)}")

    # 6. 乱码标志
    if "�" in text_all:
        problems.append("发现 U+FFFD 替换字符（乱码）")

    return problems


def main() -> None:
    # 注意：不用 Path.glob——Python 3.14 对含全角字符（项目目录名）的路径
    # glob 会静默返回空集，导致「零文件检查」误报通过；glob 模块行为正常。
    targets = [Path(a) for a in sys.argv[1:]] or sorted(map(Path, _glob.glob(str(DEFAULT_GLOB))))
    all_ok = True
    for path in targets:
        problems = check_document(path)
        if problems:
            all_ok = False
            print(f"[FAIL] {path.name}")
            for p in problems:
                print(f"       - {p}")
        else:
            print(f"[OK]   {path.name}")
    print("\n全部检查通过" if all_ok else "\n存在未通过项")
    sys.exit(0 if all_ok else 1)


if __name__ == "__main__":
    main()
