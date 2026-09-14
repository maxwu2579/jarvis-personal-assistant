"""双语学习 DOCX 渲染引擎。

排版约定（与项目文档规范一致）：
- 中文段落：Microsoft YaHei；英文段落：Aptos + 深蓝色 + 轻微缩进；
- 中英段落严格配对：英文紧跟对应中文；
- 代码：Consolas 等宽 + 浅色底纹；
- 表格：单元格内中文在上、英文在下（同一单元格上下对照）；
- 真实 Heading 样式、页眉、页码、TOC 域；
- 不使用左右双栏，不使用 Unicode 字符拼凑架构图（架构示意一律用表格）。
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "Microsoft YaHei"
EN_FONT = "Aptos"
CODE_FONT = "Consolas"
EN_COLOR = RGBColor(0x1F, 0x4E, 0x79)  # 深蓝
CODE_SHADING = "F2F2F2"
TABLE_SHADING = "F5F5F5"


def _set_font(run, *, cn=CN_FONT, en=EN_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)
    rfonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.font.italic = italic


def _set_style_east_asia(style, cn_font):
    """为样式设置中文字体（eastAsia），确保 rPr/rFonts 存在。"""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn_font)


def _shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _add_field(paragraph, instr, placeholder="（打开文档后按 F9 更新此域）"):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld_begin)
    r.append(instr_text)
    r.append(fld_sep)
    r.append(t)
    r.append(fld_end)
    return run


class BilingualDoc:
    """双语文档构建器。"""

    def __init__(self, meta: dict):
        self.doc = Document()
        self.meta = meta
        self._configure_page()
        self._configure_styles()
        self._add_header_footer()
        self._add_title_page()

    # ---------- 页面与基础样式 ----------

    def _configure_page(self):
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)

    def _configure_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = EN_FONT
        normal.font.size = Pt(11)
        _set_style_east_asia(normal, CN_FONT)

        for level, size in (("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 13)):
            style = styles[level]
            style.font.name = EN_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0x1F, 0x3B, 0x5C)
            _set_style_east_asia(style, CN_FONT)

        # 中文段落样式
        cn = styles.add_style("CnPara", WD_STYLE_TYPE.PARAGRAPH)
        cn.base_style = styles["Normal"]
        cn.font.name = EN_FONT
        _set_style_east_asia(cn, CN_FONT)
        cn.paragraph_format.space_before = Pt(12)  # 组间距
        cn.paragraph_format.space_after = Pt(2)  # 组内小间距
        cn.paragraph_format.line_spacing = 1.25

        # 英文段落样式（与中文配对，深蓝 + 缩进）
        en = styles.add_style("EnPara", WD_STYLE_TYPE.PARAGRAPH)
        en.base_style = styles["Normal"]
        en.font.name = EN_FONT
        en.font.color.rgb = EN_COLOR
        _set_style_east_asia(en, CN_FONT)
        en.paragraph_format.left_indent = Cm(0.4)
        en.paragraph_format.space_before = Pt(0)
        en.paragraph_format.space_after = Pt(12)  # 组间距
        en.paragraph_format.line_spacing = 1.25

        # 代码样式
        code = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = styles["Normal"]
        code.font.name = CODE_FONT
        code.font.size = Pt(9)
        code.paragraph_format.space_before = Pt(0)
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.line_spacing = 1.0
        _set_style_east_asia(code, CN_FONT)

        # 标签样式（如“1. 中文解释”）
        label = styles.add_style("FieldLabel", WD_STYLE_TYPE.PARAGRAPH)
        label.base_style = styles["Normal"]
        label.font.name = EN_FONT
        label.font.bold = True
        label.font.size = Pt(11)
        label.font.color.rgb = RGBColor(0x8A, 0x4B, 0x08)
        _set_style_east_asia(label, CN_FONT)
        label.paragraph_format.space_before = Pt(10)
        label.paragraph_format.space_after = Pt(2)

        # 表格文本样式
        cell = styles.add_style("CellCn", WD_STYLE_TYPE.PARAGRAPH)
        cell.base_style = styles["Normal"]
        cell.font.name = EN_FONT
        cell.font.size = Pt(9.5)
        _set_style_east_asia(cell, CN_FONT)
        cell.paragraph_format.space_before = Pt(2)
        cell.paragraph_format.space_after = Pt(2)

        cell_en = styles.add_style("CellEn", WD_STYLE_TYPE.PARAGRAPH)
        cell_en.base_style = styles["Normal"]
        cell_en.font.name = EN_FONT
        cell_en.font.size = Pt(9.5)
        cell_en.font.color.rgb = EN_COLOR
        cell_en.font.italic = True
        _set_style_east_asia(cell_en, CN_FONT)
        cell_en.paragraph_format.space_before = Pt(1)
        cell_en.paragraph_format.space_after = Pt(2)

    # ---------- 页眉页脚与封面 ----------

    def _add_header_footer(self):
        section = self.doc.sections[0]
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        run = hp.add_run(self.meta["header"])
        _set_font(run, size=Pt(9), color=RGBColor(0x80, 0x80, 0x80))
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("第 ")
        _set_font(run, size=Pt(9), color=RGBColor(0x80, 0x80, 0x80))
        _add_field(fp, "PAGE", "1")
        for r in fp.runs[1:]:
            _set_font(r, size=Pt(9), color=RGBColor(0x80, 0x80, 0x80))
        run = fp.add_run(" 页")
        _set_font(run, size=Pt(9), color=RGBColor(0x80, 0x80, 0x80))

    def _add_title_page(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(80)
        run = p.add_run(self.meta["title_cn"])
        _set_font(run, size=Pt(28), bold=True)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.meta["title_en"])
        _set_font(run, size=Pt(18), color=EN_COLOR)

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(20)
        run = p.add_run(self.meta["subtitle_cn"])
        _set_font(run, size=Pt(12), color=RGBColor(0x60, 0x60, 0x60))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(self.meta["subtitle_en"])
        _set_font(run, size=Pt(12), color=RGBColor(0x60, 0x60, 0x60))

        self.doc.add_page_break()

    # ---------- 块渲染 ----------

    def add_heading(self, text_cn, text_en, level=1):
        h = self.doc.add_heading("", level=level)
        run = h.add_run(text_cn)
        _set_font(run, size=Pt({1: 20, 2: 16, 3: 13}[level]), bold=True)
        run2 = h.add_run("  /  " + text_en)
        _set_font(run2, size=Pt({1: 20, 2: 16, 3: 13}[level]), color=EN_COLOR, bold=True)

    def add_pair(self, cn: str, en: str):
        """一对中英段落：英文紧跟中文。"""
        p = self.doc.add_paragraph(style="CnPara")
        p.add_run(cn)
        p = self.doc.add_paragraph(style="EnPara")
        p.add_run(en)

    def add_cn_only(self, text: str, style="CnPara"):
        self.doc.add_paragraph(text, style=style)

    def add_en_only(self, text: str):
        self.doc.add_paragraph(text, style="EnPara")

    def add_bullet_pair(self, cn: str, en: str | None = None):
        p = self.doc.add_paragraph(style="List Bullet")
        p.add_run(cn)
        if en:
            p = self.doc.add_paragraph(style="List Bullet")
            run = p.add_run(en)
            _set_font(run, color=EN_COLOR)

    def add_label(self, text):
        self.doc.add_paragraph(text, style="FieldLabel")

    def add_code(self, code: str, lang="python"):
        for line in code.splitlines() or [""]:
            p = self.doc.add_paragraph(style="CodeBlock")
            run = p.add_run(line)
            _set_font(run, cn=CN_FONT, en=CODE_FONT, size=Pt(9))
            _shade_paragraph(p, CODE_SHADING)
        # 代码块与周围间距
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(4)
        spacer.paragraph_format.space_after = Pt(4)

    def add_code_explain(self, block: dict):
        """完整代码讲解：代码 + 代码作用/Purpose + 逐步执行 + 为什么这样设计 + 常见错误。"""
        self.add_label(f"代码位置：{block['location']}")
        self.add_code(block["code"])
        for item in block["sections"]:
            kind = item["kind"]
            if kind == "purpose":
                self.add_label("代码作用")
                self.add_cn_only(item["cn"])
                self.add_label("Purpose")
                self.add_en_only(item["en"])
            elif kind == "steps":
                self.add_label("逐步执行")
                for cn, en in item["steps"]:
                    self.add_bullet_pair(cn, en)
                self.add_label("Step-by-step execution")
                self.add_en_only(item.get("en_intro", ""))
            elif kind == "why":
                self.add_label("为什么这样设计")
                self.add_cn_only(item["cn"])
                self.add_label("Why it is designed this way")
                self.add_en_only(item["en"])
            elif kind == "mistake":
                self.add_label("常见错误")
                self.add_cn_only(item["cn"])
                self.add_label("Common mistake")
                self.add_en_only(item["en"])

    def add_table(self, headers, rows, widths=None):
        """表格：单元格内中文在上、英文在下。

        headers: 列标题列表，每项为 "中文 / English" 混合字符串；
        rows: 每行为一个列表，每项是 (中文, English) 元组（上下对照）。
        """
        n_cols = len(headers)
        table = self.doc.add_table(rows=1 + len(rows), cols=n_cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if widths:
            for col, w in enumerate(widths):
                for cell in table.columns[col].cells:
                    cell.width = Cm(w)
        # 表头
        for col, header in enumerate(headers):
            cell = table.cell(0, col)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(header)
            _set_font(run, size=Pt(9.5), bold=True)
            self._shade_cell(cell, "DEEAF6")
        # 数据行：每项为 (中文, English) 元组 → 单元格上下对照；字符串 → 直接使用
        for r, row in enumerate(rows, start=1):
            for col, item in enumerate(row):
                cell = table.cell(r, col)
                cell.text = ""
                if isinstance(item, tuple):
                    cn, en = item
                else:
                    cn, en = str(item), None
                p = cell.paragraphs[0]
                p.style = self.doc.styles["CellCn"]
                run = p.add_run(cn)
                _set_font(run, size=Pt(9.5))
                if en:
                    p = cell.add_paragraph(style="CellEn")
                    run = p.add_run(en)
                    _set_font(run, size=Pt(9.5), italic=True, color=EN_COLOR)
                if r % 2 == 0:
                    self._shade_cell(cell, TABLE_SHADING)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_architecture_table(self, title_cn, title_en, rows):
        """架构示意表格（不用 Unicode 拼图）。rows: [(层中文, 层英文, 组件说明中文, 组件说明英文)]"""
        self.add_pair(title_cn, title_en)
        converted = [
            [(layer_cn, layer_en), (desc_cn, desc_en)]
            for layer_cn, layer_en, desc_cn, desc_en in rows
        ]
        self.add_table(
            ("层 / Layer", "说明 / Description"),
            converted,
            widths=[6.0, 10.5],
        )

    def add_toc(self):
        self.doc.add_page_break()
        h = self.doc.add_heading("", level=1)
        run = h.add_run("目录  /  Table of Contents")
        _set_font(run, size=Pt(20), bold=True)
        p = self.doc.add_paragraph()
        _add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "（目录域：在 Word 中按 F9 或右键更新）")

    def save(self, path):
        self.doc.save(path)

    @staticmethod
    def _shade_cell(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)
